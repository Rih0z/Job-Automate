#!/usr/bin/env python3
"""extract_deck.py — .pptx から事実情報（deck-facts.json）を抽出する。

使い方: python extract_deck.py <deck.pptx> -o <deck-facts.json>
依存: python-pptx。

抽出の定義:
  title        タイトルプレースホルダの文字列（strip 後。無い／空なら null）
  body_chars   タイトル以外の全テキスト（表のセル・グループ内を再帰的に含む。ノート・改行は数えない）
  bullet_count タイトル以外のテキストフレームで strip 後に非空の段落数（グループ内は再帰。表のセルは数えない）
  min_font_pt  明示指定された run のフォントサイズ(pt)の最小値。明示指定が無ければ null（継承値は取れない）
  image_count  画像（picture）シェイプ数（グループ内を再帰）
  has_notes    ノートに非空テキストがあるか
  paragraphs   タイトル以外の非空段落のテキスト（表のセルは含まない。reviewer が内容を判定するための事実情報）
limitations には `font inherited: N slides` と `font checked: M/TOTAL slides` を必ず出す。
"""
import argparse
import json
import os
import re
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

for _s in (sys.stdout, sys.stderr):
    if hasattr(_s, "reconfigure"):
        _s.reconfigure(encoding="utf-8")


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)
        else:
            yield sh


def plain(text):
    return re.sub(r"[\r\n\v\x0b]", "", text)


def number(v):
    return int(v) if float(v).is_integer() else v


def slide_facts(no, slide):
    title_shape = slide.shapes.title
    title = title_shape.text_frame.text.strip() if title_shape is not None else ""
    body_chars = bullets = images = 0
    sizes = []
    paragraphs = []
    has_text = False

    def take_runs(paragraphs):
        for p in paragraphs:
            for r in p.runs:
                if r.font.size is not None:
                    sizes.append(r.font.size.pt)

    for sh in walk(slide.shapes):
        if title_shape is not None and sh.shape_id == title_shape.shape_id:
            continue
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            images += 1
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    text = plain(cell.text_frame.text)
                    body_chars += len(text)
                    has_text = has_text or bool(text.strip())
                    take_runs(cell.text_frame.paragraphs)
        elif getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                text = plain(p.text)
                body_chars += len(text)
                if text.strip():
                    bullets += 1
                    has_text = True
                    paragraphs.append(text.strip())
            take_runs(sh.text_frame.paragraphs)
    notes = slide.has_notes_slide and bool(slide.notes_slide.notes_text_frame.text.strip())
    return {"no": no, "title": title or None, "body_chars": body_chars, "bullet_count": bullets,
            "min_font_pt": number(min(sizes)) if sizes else None, "image_count": images,
            "has_notes": bool(notes), "paragraphs": paragraphs}, has_text


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx")
    ap.add_argument("-o", "--out", required=True)
    args = ap.parse_args(argv)
    try:
        prs = Presentation(args.pptx)
    except Exception as e:  # python-pptx は形式不正で多様な例外を出す
        print("ERROR: cannot open %s: %s" % (args.pptx, e), file=sys.stderr)
        return 2
    slides, inherited = [], 0
    for i, slide in enumerate(prs.slides, 1):
        facts, has_text = slide_facts(i, slide)
        if has_text and facts["min_font_pt"] is None:
            inherited += 1
        slides.append(facts)
    checked = sum(1 for s in slides if s["min_font_pt"] is not None)
    out = {"file": os.path.basename(args.pptx), "slide_count": len(slides), "slides": slides,
           "limitations": ["font inherited: %d slides" % inherited,
                           "font checked: %d/%d slides" % (checked, len(slides))]}
    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
    return 0


if __name__ == "__main__":
    sys.exit(main())
